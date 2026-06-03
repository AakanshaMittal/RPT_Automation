import os
import csv
import re
import json
from pathlib import Path
from typing import List, Dict
from html import escape

from openpyxl import Workbook
from openpyxl.styles import Font

try:
    import pandas as pd
except Exception:
    pd = None


BASE_PATH = Path("data/ratesPerQuantity")
FILE1_FOLDER = BASE_PATH / "File1"
FILE2_FOLDER = BASE_PATH / "File2"
OUTPUT_FOLDER = BASE_PATH / "Output"
mapping_file = r"config/mapping.json"


# -------------------- MAPPING METHODS --------------------

def load_mapping():
    with open(mapping_file) as f:
        return json.load(f)


def reverse_mapping(mapping):
    return {v: k for k, v in mapping.items()}


def extract_rpt_id(filename):
    match = re.search(r"(RPT\d+[A-Z]{2})", filename.upper())
    return match.group(1) if match else None


def extract_tc(filename):
    match = re.search(r"TC(\d+)", filename.upper())
    return match.group(1) if match else None


# -------------------- UTILITIES --------------------

def clean_text(text: str) -> str:
    if not text:
        return ""
    return str(text).replace("\ufeff", "").strip()


def normalize_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", clean_text(h).lower())


def clean_numeric(val: str) -> str:
    if not val:
        return ""
    return val.replace("$", "").replace(",", "").replace("%", "").strip()


def is_number(val: str) -> bool:
    try:
        float(clean_numeric(val))
        return True
    except Exception:
        return False


def numeric_major_mismatch(v1: str, v2: str) -> bool:
    try:
        return abs(float(clean_numeric(v1)) - float(clean_numeric(v2))) > 0.0001
    except Exception:
        return False


def normalize_ndc(val: str) -> str:
    val = clean_text(val)
    if not val:
        return ""
    if val.isdigit():
        return val.zfill(11)
    return val


def extract_year(text: str) -> str:
    m = re.search(r"(20\d{2})", text)
    return m.group(1) if m else ""


def read_excel_safe(path: str) -> List[List[str]]:
    if pd is None:
        raise RuntimeError("pandas not installed")

    df = pd.read_excel(path, dtype=str)
    df = df.fillna("")
    df = df.loc[:, (df != "").any(axis=0)]
    return [[str(c).strip() for c in row] for row in df.values.tolist()]


def read_csv_safe(path: str, delimiter: str):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter=delimiter))
        return [r[1:] if len(r) > 1 else r for r in rows]


def read_table(path: str, delimiter: str):
    ext = os.path.splitext(path)[1].lower()
    return read_excel_safe(path) if ext in (".xls", ".xlsx") else read_csv_safe(path, delimiter)


# -------------------- HEADER HANDLING --------------------

def build_composite_headers(rows, group_row, detail_row):
    groups = rows[group_row - 1]
    details = rows[detail_row - 1]

    max_cols = max(len(groups), len(details))
    groups += [""] * (max_cols - len(groups))
    details += [""] * (max_cols - len(details))

    headers = []
    current_group = ""

    for i in range(max_cols):
        g = clean_text(groups[i])
        d = clean_text(details[i])

        if g:
            current_group = g

        if not d:
            headers.append("")
            continue

        raw = f"{current_group}_{d}" if current_group else d
        year = extract_year(raw)
        headers.append(f"{current_group}_YEAR_{year}" if year else raw)

    return headers


def build_index(headers: List[str]) -> Dict[str, int]:
    return {h: i for i, h in enumerate(headers) if h}


def find_missing_columns(headers1, headers2):
    h1 = {h for h in headers1 if h}
    h2 = {h for h in headers2 if h}
    return sorted(h2 - h1), sorted(h1 - h2)


def build_row_map(rows, headers, data_row):
    norm_map = {normalize_header(h): i for i, h in enumerate(headers) if h}

    ndc_col = next((i for h, i in norm_map.items() if "ndc" in h), None)
    if ndc_col is None:
        raise ValueError("NDC column not found")

    row_map = {}
    for r in rows[data_row - 1:]:
        ndc = normalize_ndc(r[ndc_col] if ndc_col < len(r) else "")
        if ndc:
            row_map[ndc] = r

    return row_map


def extract_group(header: str):
    return header.split("_")[0] if "_" in header else header


def compare_by_row_key(headers1, headers2, map1, map2):
    common_headers = sorted(h for h in (set(headers1) & set(headers2)) if h)
    idx1 = build_index(headers1)
    idx2 = build_index(headers2)

    mismatches = []
    group_counter = {}

    extra1 = sorted(set(map1) - set(map2))
    extra2 = sorted(set(map2) - set(map1))

    common_keys = set(map1) & set(map2)

    for key in common_keys:
        r1 = map1[key]
        r2 = map2[key]

        for h in common_headers:
            if h not in idx1 or h not in idx2:
                continue

            v1 = r1[idx1[h]] if idx1[h] < len(r1) else ""
            v2 = r2[idx2[h]] if idx2[h] < len(r2) else ""

            v1 = v1.strip()
            v2 = v2.strip()

            mismatch = False
            highlight = False

            if is_number(v1) and is_number(v2):
                if numeric_major_mismatch(v1, v2):
                    mismatch = True
                    highlight = True
            elif v1 != v2:
                mismatch = True

            if mismatch:
                mismatches.append((key, h, v1, v2, highlight))
                g = extract_group(h)
                group_counter[g] = group_counter.get(g, 0) + 1

    return mismatches, group_counter, extra1, extra2


# -------------------- HTML REPORT (UNCHANGED) --------------------

def write_html_report(path, mismatches, group_counter, extra1, extra2,
                      file1, file2, missing1, missing2):

    def li(items):
        return "".join(f"<li>{escape(i)}</li>" for i in items) or "<li>None</li>"

    mis_rows = ""
    for k_toggle, h, v1, v2, hi in mismatches:
        style = "background:#ffd6d6;font-weight:600" if hi else ""
        mis_rows += (
            f"<tr style='{style}'>"
            f"<td>{k_toggle}</td><td>{h}</td>"
            f"<td>{escape(v1)}</td><td>{escape(v2)}</td></tr>"
        )

    html = f"""
<html>
<head>
<title>Data Validation Report</title>
<style>
body{{font-family:Segoe UI;margin:20px}}
table{{border-collapse:collapse;width:100%;margin-bottom:30px}}
th,td{{border:1px solid #ccc;padding:6px}}
th{{background:#eee}}
</style>
</head>
<body>

<h2>Data Validation Report</h2>
<p><b>Total mismatches:</b> {len(mismatches)}</p>

<h3>Missing Columns</h3>
<b>Missing in {os.path.basename(file1)}</b>
<ul>{li(missing1)}</ul>

<b>Missing in {os.path.basename(file2)}</b>
<ul>{li(missing2)}</ul>

<h3>Extra Rows</h3>
<b>Extra in {os.path.basename(file1)}</b>
<ul>{li(extra1)}</ul>

<b>Extra in {os.path.basename(file2)}</b>
<ul>{li(extra2)}</ul>

<h3>Group Summary</h3>
<table>
<tr><th>Group</th><th>Count</th></tr>
{''.join(f'<tr><td>{g}</td><td>{c}</td></tr>' for g, c in group_counter.items())}
</table>

<h3>Mismatch Details</h3>
<table>
<tr><th>NDC11</th><th>Column</th><th>{file1}</th><th>{file2}</th></tr>
{mis_rows}
</table>

</body>
</html>
"""

    with open(path, "w", encoding="utf-8") as f:
        f.write(html)


# -------------------- EXCEL WRITER --------------------

def write_tc_excel_sheet(ws, tc, mismatches, group_counter,
                          extra1, extra2, missing1, missing2,
                          file1, file2):

    bold = Font(bold=True)
    row = 1

    def title(text):
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = bold
        row += 1

    title("Data Validation Report")
    ws.append(["Total mismatches", len(mismatches)])
    row += 1

    title("Missing Columns")
    ws.append([f"Missing in {file1}"])
    for c in missing1 or ["None"]:
        ws.append([c])

    ws.append([f"Missing in {file2}"])
    for c in missing2 or ["None"]:
        ws.append([c])

    row += 1
    title("Extra Rows")

    ws.append([f"Extra in {file1}"])
    for r_val in extra1 or ["None"]:
        ws.append([r_val])

    ws.append([f"Extra in {file2}"])
    for r_val in extra2 or ["None"]:
        ws.append([r_val])

    row += 1
    title("Group Summary")
    ws.append(["Group", "Count"])
    for g, c in group_counter.items():
        ws.append([g, c])

    row += 1
    title("Mismatch Details")
    ws.append(["NDC11", "Column", file1, file2, "Highlighted"])
    for k_toggle, h, v1, v2, hi in mismatches:
        ws.append([k_toggle, h, v1, v2, "YES" if hi else "NO"])


# -------------------- MAIN --------------------

def main():
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

    mapping = load_mapping()
    rpt_to_tc = reverse_mapping(mapping)

    file1_map = {}
    file2_map = {}

    for f in FILE1_FOLDER.glob("*.xlsx"):
        rpt_id = extract_rpt_id(f.name)
        if rpt_id and rpt_id in rpt_to_tc:
            file1_map[rpt_to_tc[rpt_id].replace("TC", "")] = f

    for f in FILE2_FOLDER.glob("*.xlsx"):
        tc = extract_tc(f.name)
        if tc:
            file2_map[tc] = f

    common_tcs = sorted(set(file1_map) & set(file2_map), key=int)

    workbook = Workbook()
    ws_summary = workbook.active
    ws_summary.title = "Summary"
    ws_summary.append(["Test Case", "Total Mismatches"])

    for tc in common_tcs:
        file1 = file1_map[tc]
        file2 = file2_map[tc]

        rows1 = read_table(file1, ",")
        rows2 = read_table(file2, ",")

        headers1 = build_composite_headers(rows1, 1, 2)
        headers2 = build_composite_headers(rows2, 1, 2)

        missing1, missing2 = find_missing_columns(headers1, headers2)

        map1 = build_row_map(rows1, headers1, 3)
        map2 = build_row_map(rows2, headers2, 3)

        mismatches, group_counter, extra1, extra2 = compare_by_row_key(
            headers1, headers2, map1, map2
        )

        ws_summary.append([f"TC{tc}", len(mismatches)])

        ws_tc = workbook.create_sheet(f"TC{tc}")
        write_tc_excel_sheet(
            ws_tc, tc, mismatches, group_counter,
            extra1, extra2, missing1, missing2,
            file1.name, file2.name
        )

        write_html_report(
            OUTPUT_FOLDER / f"TC{tc}_RPQ_Report_Dev3Jun.html",
            mismatches, group_counter, extra1, extra2,
            file1.name, file2.name, missing1, missing2
        )

    workbook.save(OUTPUT_FOLDER / "Consolidated_RPQ_Report_Dev3Jun.xlsx")
    print("All scenarios completed.")


if __name__ == "__main__":
    main()