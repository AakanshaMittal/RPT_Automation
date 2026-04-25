import pandas as pd
import re
import json
from pathlib import Path
from collections import defaultdict
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font

file1_folder = r"data/raOutput/File1"
file2_folder = r"data/raOutput/File2"
output_folder = r"data/raOutput/Output"
mapping_file = r"config/mapping.json"

TOL = "0.0001"

start_col_file1 = 0
start_col_file2 = 0


def load_mapping():
    with open(mapping_file) as f:
        return json.load(f)


def reverse_mapping(mapping):
    return {v: k for k, v in mapping.items()}


def extract_rpt_id(filename):
    match = re.search(r"(RPT\d+[A-Z]{2})", filename.upper())
    return match.group(1) if match else None


def extract_tc(filename):
    match = re.search(r"TC(\d+)", filename.upper())
    return match.group(1) if match else None


def normalize_year(text):
    if pd.isna(text):
        return ""
    t = str(text)
    cot = " CoT" if "cot" in t.lower() else ""
    m = re.search(r"(20\d{2})", t)
    if m:
        return m.group(1) + cot
    m = re.search(r"01/(\d{2})-12/(\d{2})", t)
    if m:
        return "20" + m.group(2) + cot
    return t.strip()


def clean_text(x):
    if pd.isna(x):
        return ""
    x = str(x)
    x = x.replace('"', '').replace('“', '').replace('”', '')
    return re.sub(r"\s+", " ", x).strip().upper()


def base_header(col):
    return re.sub(r'_(20\d{2})(\s*COT)?$', '', col, flags=re.I)


def build_headers(df, header_row):
    raw_h1 = df.iloc[header_row]
    raw_h2 = df.iloc[header_row + 1]
    raw_h3 = df.iloc[header_row + 2]

    h1, h2 = raw_h1.copy(), raw_h2.copy()
    last_h1, last_h2 = "", ""

    for i in range(len(df.columns)):
        c1 = "" if pd.isna(raw_h1.iloc[i]) else str(raw_h1.iloc[i]).strip()
        c2 = "" if pd.isna(raw_h2.iloc[i]) else str(raw_h2.iloc[i]).strip()
        c3 = "" if pd.isna(raw_h3.iloc[i]) else str(raw_h3.iloc[i]).strip()

        if c1 == "" and c2 == "" and c3 == "":
            h1.iloc[i] = ""
            h2.iloc[i] = ""
            continue

        if c1 == "" and c2 and c3:
            c1 = last_h1

        if c1:
            last_h1, last_h2 = c1, ""
        if c2:
            last_h2 = c2

        h1.iloc[i], h2.iloc[i] = last_h1, last_h2

    final_cols = []
    for i in range(len(df.columns)):
        p1 = str(h1.iloc[i]).strip()
        p2 = str(h2.iloc[i]).strip()
        p3 = normalize_year(str(raw_h3.iloc[i]).strip())
        col = "_".join(p for p in [p1, p2, p3] if p and p.lower() != "nan")
        col = re.sub(r"_+", "_", col)
        final_cols.append(col if col else f"EMPTY_COL_{i}")

    data = df.iloc[header_row + 3:].copy()
    data.columns = final_cols
    data = data.loc[:, ~data.columns.duplicated()]
    return data.reset_index(drop=True)


def load_file(path, header_row, start_col):
    if str(path).lower().endswith(".csv"):
        raw = pd.read_csv(path, header=None, dtype=str)
    else:
        raw = pd.read_excel(path, header=None, dtype=str)

    raw = raw.iloc[:, start_col:].reset_index(drop=True)
    df = build_headers(raw, header_row)
    return df.map(clean_text)


def normalize_ndc_value(x):
    if x is None:
        return ""
    s = str(x).strip()
    return s.lstrip("0") if s.isdigit() else s


def normalize_ndc_columns(df):
    for c in [c for c in df.columns if "NDC" in c.upper()]:
        df[c] = df[c].apply(normalize_ndc_value)
    return df


def make_row_key(df, n=10):
    df = normalize_ndc_columns(df)
    key_cols = df.columns[:n]
    df["__ROW_KEY__"] = df[key_cols].astype(str).agg("||".join, axis=1)
    return df


def safe_value(x):
    if isinstance(x, pd.Series):
        return x.iloc[0] if len(x) else ""
    return x


def to_decimal(x):
    if x is None:
        return None
    v = str(x).replace("$", "").replace(",", "").strip()
    if v == "":
        return None
    try:
        return Decimal(v)
    except Exception:
        return None


def values_equal(v1, v2):
    d1, d2 = to_decimal(v1), to_decimal(v2)
    if d1 is not None and d2 is not None:
        return abs(d1 - d2) <= Decimal(TOL)
    return str(v1) == str(v2)


def needs_highlight(v1, v2):
    d1, d2 = to_decimal(v1), to_decimal(v2)
    return d1 is not None and d2 is not None and abs(d1 - d2) > Decimal(TOL)


def normalize_col_tokens(col):
    col = base_header(col)
    col = re.sub(r'[^A-Z0-9 ]', ' ', col.upper())
    return re.sub(r'\s+', ' ', col).strip()


def tail_signature(col, n=3):
    tokens = normalize_col_tokens(col).split()
    return " ".join(tokens[-n:]) if len(tokens) >= n else " ".join(tokens)


def build_column_mapping(cols1, cols2):
    map12, used = {}, set()
    norm2 = {c: normalize_col_tokens(c) for c in cols2}

    for c1 in cols1:
        n1 = normalize_col_tokens(c1)
        strict = [c2 for c2, n2 in norm2.items() if n1 == n2 and c2 not in used]
        if strict:
            map12[c1] = strict[0]
            used.add(strict[0])
            continue

        t1 = tail_signature(c1)
        fallback = [c2 for c2 in cols2 if tail_signature(c2) == t1 and c2 not in used]
        if len(fallback) == 1:
            map12[c1] = fallback[0]
            used.add(fallback[0])

    return map12


def write_tc_excel_sheet(ws, tc, mismatches, summary, ndc_map, extra1, extra2, file1, file2):
    bold = Font(bold=True)
    row = 1

    def title(text):
        nonlocal row
        ws.cell(row=row, column=1, value=text).font = bold
        row += 1

    title("RA Output Validation Report")
    ws.append(["Total mismatches", len(mismatches)])
    row += 1

    ws.append([f"Extra rows in {file1}", len(extra1)])
    ws.append([f"Extra rows in {file2}", len(extra2)])
    row += 2

    title("Mismatch Summary (with NDCs)")
    ws.append(["Header", "Mismatch Count", "Affected NDCs"])
    for k in sorted(summary, key=lambda x: -summary[x]):
        ws.append([k, summary[k], ", ".join(sorted(ndc_map[k]))])

    row += 2

    title("Value Mismatches")
    ws.append(["Row Key", "Column", "File 1", "File 2", "Highlighted"])
    for rk, col, v1, v2, hl in mismatches:
        ws.append([rk, col, v1, v2, "YES" if hl else "NO"])

    row += 2

    title("Extra Rows - Only in File 1")
    for _, r in extra1.iterrows():
        ws.append(r.tolist())

    row += 1
    title("Extra Rows - Only in File 2")
    for _, r in extra2.iterrows():
        ws.append(r.tolist())


def generate_html_report(extra1, extra2, mismatches, summary, ndc_map, file1, file2, out_path):
    html = []
    html.append("<html><head><title>RA Output Validation</title>")
    html.append("<style>")
    html.append("body{font-family:Arial;padding:20px}")
    html.append("table{border-collapse:collapse;width:100%}")
    html.append("th,td{border:1px solid #ccc;padding:6px;font-size:12px}")
    html.append("th{background:#f2f2f2}")
    html.append(".diff{background:#ffd6d6;font-weight:600}")
    html.append("</style></head><body>")
    html.append("<h1>RA Output Validation Report</h1>")
    html.append(f"<p><b>Total mismatches:</b> {len(mismatches)}</p>")
    html.append(f"<p><b>Extra rows in {file1}:</b> {len(extra1)}</p>")
    html.append(f"<p><b>Extra rows in {file2}:</b> {len(extra2)}</p>")

    html.append("<h2>Mismatch Summary (with NDCs)</h2>")
    html.append("<table><tr><th>Header</th><th>Mismatch Count</th><th>Affected NDCs</th></tr>")
    for k, v in sorted(summary.items(), key=lambda x: -x[1]):
        html.append(f"<tr><td>{k}</td><td>{v}</td><td>{', '.join(sorted(ndc_map[k]))}</td></tr>")
    html.append("</table>")

    html.append("<h2>Value Mismatches</h2>")
    html.append("<table><tr><th>Row Key</th><th>Column</th><th>File 1</th><th>File 2</th></tr>")
    for rk, col, v1, v2, hl in mismatches:
        cls = "diff" if hl else ""
        html.append(f"<tr class='{cls}'><td>{rk}</td><td>{col}</td><td>{v1}</td><td>{v2}</td></tr>")
    html.append("</table>")

    html.append("<h2>Extra Rows</h2><h3>Only in File 1</h3>")
    html.append(extra1.to_html(index=False))
    html.append("<h3>Only in File 2</h3>")
    html.append(extra2.to_html(index=False))
    html.append("</body></html>")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(html))


def main():
    h1, h2 = 0, 6

    out_path = Path(output_folder)
    out_path.mkdir(parents=True, exist_ok=True)

    mapping = load_mapping()
    rpt_to_tc = reverse_mapping(mapping)

    file1_map, file2_map = {}, {}

    for f in Path(file1_folder).glob("*.xlsx"):
        rpt = extract_rpt_id(f.name)
        if rpt and rpt in rpt_to_tc:
            file1_map[rpt_to_tc[rpt].replace("TC", "")] = f

    for f in Path(file2_folder).glob("TC*_Exc_RAO*.xlsx"):
        tc = extract_tc(f.name)
        if tc:
            file2_map[tc] = f

    workbook = Workbook()
    ws_summary = workbook.active
    ws_summary.title = "Summary"
    ws_summary.append(["Test Case", "Total Mismatches"])

    for tc in sorted(set(file1_map) & set(file2_map), key=int):
        file1 = file1_map[tc]
        file2 = file2_map[tc]

        df1 = make_row_key(load_file(file1, h1, start_col_file1), 10).set_index("__ROW_KEY__")
        df2 = make_row_key(load_file(file2, h2, start_col_file2), 10).set_index("__ROW_KEY__")

        col_map = build_column_mapping(df1.columns, df2.columns)

        extra1 = df1.loc[~df1.index.isin(df2.index)].reset_index()
        extra2 = df2.loc[~df2.index.isin(df1.index)].reset_index()

        mismatches = []
        summary = defaultdict(int)
        ndc_map = defaultdict(set)

        for key in df1.index.intersection(df2.index):
            ndc = key.split("||")[1] if "||" in key else key
            r1, r2 = df1.loc[key], df2.loc[key]

            for c1, c2 in col_map.items():
                v1 = safe_value(r1.get(c1, ""))
                v2 = safe_value(r2.get(c2, ""))

                if not values_equal(v1, v2):
                    hl = needs_highlight(v1, v2)
                    header = base_header(c1)
                    mismatches.append((key, header, v1, v2, hl))
                    summary[header] += 1
                    ndc_map[header].add(ndc)

        ws_summary.append([f"TC{tc}", len(mismatches)])

        ws_tc = workbook.create_sheet(f"TC{tc}")
        write_tc_excel_sheet(ws_tc, tc, mismatches, summary, ndc_map, extra1, extra2, file1.name, file2.name)

        generate_html_report(
            extra1, extra2, mismatches, summary, ndc_map,
            file1.name, file2.name,
            out_path / f"TC{tc}_RAO_Report_Dev24Apr.html"
        )

    workbook.save(out_path / "Consolidated_RAO_Report_Dev24Apr.xlsx")
    print("All scenarios completed.")


if __name__ == "__main__":
    main()