import pandas as pd

from html import escape

import os

import json

import re

from pathlib import Path

from openpyxl.styles import PatternFill
 
BASE_PATH = Path("data/externalStats")
 
FILE1_DIR = BASE_PATH / "File1"

FILE2_DIR = BASE_PATH / "File2"

OUTPUT_DIR = BASE_PATH / "Output"

MAPPING_FILE = "config/mapping.json"
 
os.makedirs(OUTPUT_DIR, exist_ok=True)
 
 
# ------------------ UTIL ------------------

def load_mapping():

    with open(MAPPING_FILE) as f:

        return json.load(f)
 
 
def reverse_mapping(mapping):

    return {v: k for k, v in mapping.items()}
 
 
def extract_rpt_id(filename):

    match = re.search(r"(RPT\d+[A-Z]{2})", filename.upper())

    return match.group(1) if match else None
 
 
def extract_tc(filename):

    match = re.search(r"TC(\d+)", filename.upper())

    return f"TC{match.group(1)}" if match else None
 
 
def clean(x):

    return "" if pd.isna(x) else str(x).strip()
 
 
def try_float(x):

    try:

        return float(str(x).replace(",", "").replace("%", ""))

    except:

        return None
 
 
def trim_table(df):

    df = df.dropna(how="all")

    df = df.loc[:, df.notna().any()]

    return df.reset_index(drop=True)
 
 
# ------------------ EXCEL HELPERS ------------------

def excel_col_to_index(col):

    index = 0

    for c in col:

        index = index * 26 + (ord(c.upper()) - ord('A') + 1)

    return index - 1
 
 
def get_range(df, cell_range):

    start, end = cell_range.split("-")
 
    col1 = re.findall(r"[A-Z]+", start)[0]

    row1 = int(re.findall(r"\d+", start)[0])
 
    col2 = re.findall(r"[A-Z]+", end)[0]

    row2 = int(re.findall(r"\d+", end)[0])
 
    c1 = excel_col_to_index(col1)

    c2 = excel_col_to_index(col2)
 
    return df.iloc[row1 - 1:row2, c1:c2 + 1]
 
 
# ------------------ TABLE NAME ------------------

def extract_table_name(table):

    for i in range(min(5, len(table))):

        val = clean(table.iloc[i, 0])

        if val:

            return val.lower()

    return "table"
 
 
# ------------------ FIXED TABLE EXTRACTION ------------------

TABLE_RANGES = [

    "B23-H30",

    "B32-H38",

    "B40-H46",

    "B51-Z68",

    "B70-Z86",

    "B88-Z108",

    "AE23-AI44",

    "AK23-AN44",

    "AE47-AH68"

]
 
 
def extract_tables(df):

    tables = {}
 
    for rng in TABLE_RANGES:

        table = trim_table(get_range(df, rng).copy())
 
        if table.shape[0] < 2:

            continue
 
        name = extract_table_name(table)
 
        key = name

        if key in tables:

            key = f"{key}_{len(tables)}"
 
        tables[key] = table.reset_index(drop=True)
 
    return tables
 
 
# ------------------ ALIGN ------------------

def get_drug_name_column(df):

    for col in range(df.shape[1]):

        for row in range(min(3, df.shape[0])):

            if "drug name" in str(df.iloc[row, col]).lower():

                return col

    return None
 
 
def align_tables(t1, t2):

    k1, k2 = get_drug_name_column(t1), get_drug_name_column(t2)

    if k1 is None or k2 is None:

        return t1, t2
 
    def extract_data(df, c):

        data, header = {}, 0
 
        for i in range(min(3, len(df))):

            if "drug name" in str(df.iloc[i, c]).lower():

                header = i

                break
 
        for i in range(header + 1, len(df)):

            drug = clean(df.iloc[i, c])

            if drug:

                data[drug] = [clean(df.iloc[i, j]) for j in range(df.shape[1])]
 
        return data
 
    d1, d2 = extract_data(t1, k1), extract_data(t2, k2)

    all_keys = sorted(set(d1) | set(d2))
 
    max_cols = max(t1.shape[1], t2.shape[1])
 
    a1, a2 = [], []

    for k in all_keys:

        r1 = d1.get(k, [""] * max_cols)

        r2 = d2.get(k, [""] * max_cols)
 
        a1.append(r1 + [""] * (max_cols - len(r1)))

        a2.append(r2 + [""] * (max_cols - len(r2)))
 
    return pd.DataFrame(a1), pd.DataFrame(a2)
 
 
# ------------------ COMPARE ------------------

def compare_tables(t1, t2, tol=1e-6):

    mismatches = []
 
    for i in range(min(len(t1), len(t2))):

        for j in range(min(t1.shape[1], t2.shape[1])):
 
            v1, v2 = clean(t1.iloc[i, j]), clean(t2.iloc[i, j])
 
            f1, f2 = try_float(v1), try_float(v2)
 
            if f1 is not None and f2 is not None:

                if abs(f1 - f2) > tol:

                    mismatches.append((i, j))

            elif v1 != v2:

                mismatches.append((i, j))
 
    return mismatches
 
 
# ------------------ HTML REPORT (UNCHANGED) ------------------

def generate_html(results, f1_name, f2_name, output_path):

    html = """
<html><head><style>

body {font-family: Arial;}

table {border-collapse: collapse; margin-bottom:20px;}

td, th {border:1px solid black; padding:6px;}

.mismatch {background:#ffcccc;}

.flex {display:flex; gap:40px;}
</style></head><body>
<h1>Detailed Stats</h1>

"""
 
    for name, (t1, t2, mm) in results.items():

        if not mm:

            continue
 
        html += f"<h3>{name}</h3><div class='flex'>"
 
        for df, label in [(t1, f1_name), (t2, f2_name)]:

            html += f"<div><h4>{label}</h4><table>"
 
            for i in range(len(df)):

                html += "<tr>"

                for j in range(df.shape[1]):

                    cls = "mismatch" if (i, j) in mm else ""

                    html += f"<td class='{cls}'>{escape(clean(df.iloc[i, j]))}</td>"

                html += "</tr>"
 
            html += "</table></div>"
 
        html += "</div>"
 
    html += "</body></html>"
 
    with open(output_path, "w", encoding="utf-8") as f:

        f.write(html)
 
 
# ------------------ CONSOLIDATED EXCEL ------------------

def generate_consolidated_excel(summary, all_results, output_path):

    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
 
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
 
        # Summary

        summary_df = pd.DataFrame(summary).fillna(0).astype(int)

        summary_df.index.name = "Table"

        summary_df.reset_index(inplace=True)

        summary_df.to_excel(writer, sheet_name="Summary", index=False)
 
        # TC Sheets

        for tc, results in all_results.items():

            sheet_name = tc[:31]
 
            ws = writer.book.create_sheet(sheet_name)

            writer.sheets[sheet_name] = ws
 
            current_row = 0
 
            for name, (t1, t2, mismatches) in results.items():
 
                if t1.empty or t2.empty or not mismatches :

                    continue
 
                ws.cell(row=current_row + 1, column=1, value=name)
 
                t1.to_excel(writer, sheet_name=sheet_name,

                            startrow=current_row + 2, startcol=0,

                            index=False, header=False)
 
                offset = t1.shape[1] + 3
 
                t2.to_excel(writer, sheet_name=sheet_name,

                            startrow=current_row + 2, startcol=offset,

                            index=False, header=False)
 
                # Highlight mismatches

                for i, j in mismatches:

                    ws.cell(row=current_row + 3 + i, column=1 + j).fill = red_fill

                    ws.cell(row=current_row + 3 + i, column=offset + 1 + j).fill = red_fill
 
                current_row += max(len(t1), 3) + 5
 
        if "Sheet" in writer.book.sheetnames:

            del writer.book["Sheet"]
 
 
# ------------------ MAIN ------------------

def main():

    mapping = load_mapping()

    rev_map = reverse_mapping(mapping)
 
    file1_dict, file2_dict = {}, {}
 
    for f in FILE1_DIR.glob("*.xlsx"):

        rpt = extract_rpt_id(f.name)

        if rpt and rpt in rev_map:

            file1_dict[rev_map[rpt]] = f
 
    for f in FILE2_DIR.glob("*.xlsx"):

        tc = extract_tc(f.name)

        if tc:

            file2_dict[tc] = f
 
    summary = {}

    all_results = {}
 
    for tc in sorted(set(file1_dict) & set(file2_dict)):

        print(f"\nProcessing {tc}...")
 
        df1 = pd.read_excel(file1_dict[tc], header=None)

        df2 = pd.read_excel(file2_dict[tc], header=None)
 
        t1_map = extract_tables(df1)

        t2_map = extract_tables(df2)
 
        results = {}

        summary[tc] = {}
 
        for key in set(t1_map) & set(t2_map):

            t1, t2 = align_tables(t1_map[key], t2_map[key])

            mm = compare_tables(t1, t2)

            if mm:
                results[key] = (t1, t2, mm)
                summary[tc][key] = len(mm)
 
        if results:

            all_results[tc] = results
 
            generate_html(

                results,

                file1_dict[tc].name,

                file2_dict[tc].name,

                OUTPUT_DIR / f"{tc}_External.html"

            )
 
    generate_consolidated_excel(

        summary,

        all_results,

        OUTPUT_DIR / "Consolidated_External_Report.xlsx"

    )
 
    print("\nAll scenarios completed.")
 
 
if __name__ == "__main__":

    main()
 